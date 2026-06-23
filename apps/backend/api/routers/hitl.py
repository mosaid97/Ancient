"""HITL review API router — OCR & translation quality review queue."""
from __future__ import annotations

import io
import logging
import os
import subprocess
import sys
from pathlib import Path
from typing import Any
from urllib.parse import quote

from fastapi import APIRouter, Body, Depends, HTTPException, Query
from fastapi.responses import Response
from neo4j import Driver

from apps.backend.api.deps import get_driver
from apps.backend.feedback.active_learning import get_review_queue
from apps.backend.feedback.correction_writer import CorrectionInput, write_correction

log = logging.getLogger(__name__)
router = APIRouter()

_REPO_ROOT = Path(__file__).resolve().parents[4]
_INDEX_RUNNER = _REPO_ROOT / "scripts" / "run_index_pipeline.py"


# ---------------------------------------------------------------------------
# Review queue
# ---------------------------------------------------------------------------


@router.get("/queue")
async def review_queue(
    max: int = Query(30, ge=1, le=200),
    driver: Driver = Depends(get_driver),
) -> dict[str, Any]:
    try:
        items = get_review_queue(driver, max_pages=max)
    except Exception as exc:
        log.exception("Failed to fetch review queue: %s", exc)
        raise HTTPException(status_code=500, detail=str(exc))

    serialized = []
    for item in items:
        d = item.to_dict() if hasattr(item, "to_dict") else (item.__dict__ if hasattr(item, "__dict__") else dict(item))
        # Alias for frontend: decision = evaluation_decision
        d["decision"] = d.get("evaluation_decision", "needs_review")
        serialized.append(d)
    return {"items": serialized, "total": len(serialized)}


# ---------------------------------------------------------------------------
# Page detail
# ---------------------------------------------------------------------------


def _fetch_page_detail(driver: Driver, page_id: str) -> dict[str, Any]:
    with driver.session() as s:
        page_row = s.run(
            """
            MATCH (p:PAGE {id: $pid})
            OPTIONAL MATCH (doc:DOCUMENT)-[:CONSIST_OF]->(:CHAPTER)-[:INCLUDE]->(:SECTION)-[:INCLUDE]->(p)
            OPTIONAL MATCH (doc2:DOCUMENT)-[:CONSIST_OF]->(:CHAPTER)-[:INCLUDE]->(p)
            WITH p,
                 coalesce(doc.title, doc2.title) AS docTitle,
                 coalesce(doc.tier, doc2.tier)   AS docTier
            RETURN
                p.id                    AS page_id,
                p.pageIndex             AS page_index,
                p.mode                  AS mode,
                p.language              AS language,
                p.tier                  AS tier,
                p.fusionStatus          AS fusion_status,
                p.textFused             AS text_fused,
                p.paddleOcrText         AS paddle_text,
                p.qwenVlOcrText         AS qwen_text,
                p.deepseekOcrText       AS deepseek_text,
                p.paddleOcrConfidence   AS paddle_confidence,
                p.qwenVlOcrConfidence   AS qwen_confidence,
                p.deepseekOcrConfidence AS deepseek_confidence,
                p.evaluationDecision    AS eval_decision,
                p.problemClass          AS problem_class,
                p.problemClassConfidence AS problem_class_confidence,
                p.problemClassReasoning  AS reasoning,
                p.interEngineCer        AS cer,
                p.cjkValidityRatio      AS cjk_ratio,
                p.evaluatedAt           AS evaluated_at,
                p.imageUri              AS image_uri,
                docTitle                AS document_title,
                docTier                 AS document_tier
            """,
            pid=page_id,
        ).single()

        if page_row is None:
            return {}

        chunks_rows = s.run(
            """
            MATCH (p:PAGE {id: $pid})-[:HAS]->(c:CHUNK)
            RETURN
                c.id              AS chunk_id,
                c.chunkIndex      AS index,
                c.text            AS text,
                c.textCanonical   AS text_canonical,
                c.textVernacular  AS text_vernacular,
                c.translationStatus AS translation_status,
                c.tier            AS tier
            ORDER BY c.chunkIndex
            """,
            pid=page_id,
        ).data()

    page = dict(page_row)
    ocr = {
        "paddle": page.pop("paddle_text", None),
        "qwen_vl": page.pop("qwen_text", None),
        "deepseek": page.pop("deepseek_text", None),
        "fused": page.pop("text_fused", None),
        "paddle_confidence": page.pop("paddle_confidence", None),
        "qwen_confidence": page.pop("qwen_confidence", None),
        "deepseek_confidence": page.pop("deepseek_confidence", None),
    }
    evaluation = {
        "decision": page.pop("eval_decision", None),
        "problem_class": page.pop("problem_class", None),
        "problem_class_confidence": page.pop("problem_class_confidence", None),
        "reasoning": page.pop("reasoning", None),
        "inter_engine_cer": page.pop("cer", None),
        "cjk_validity_ratio": page.pop("cjk_ratio", None),
        "evaluated_at": page.pop("evaluated_at", None),
    }
    return {
        "page": page,
        "ocr": ocr,
        "evaluation": evaluation,
        "chunks": chunks_rows,
    }


@router.get("/page/{page_id}")
async def page_detail(
    page_id: str,
    driver: Driver = Depends(get_driver),
) -> dict[str, Any]:
    try:
        detail = _fetch_page_detail(driver, page_id)
    except Exception as exc:
        log.exception("Page detail failed for %s: %s", page_id, exc)
        raise HTTPException(status_code=500, detail=str(exc))
    if not detail:
        raise HTTPException(status_code=404, detail=f"Page {page_id!r} not found")
    return detail


# ---------------------------------------------------------------------------
# Actions
# ---------------------------------------------------------------------------


@router.post("/approve")
async def approve_page(
    payload: dict = Body(...),
    driver: Driver = Depends(get_driver),
) -> dict[str, Any]:
    page_id = payload.get("page_id")
    if not page_id:
        raise HTTPException(status_code=422, detail="page_id required")
    try:
        with driver.session() as s:
            s.run(
                """
                MATCH (p:PAGE {id: $pid})
                SET p.hitlStatus = 'approved',
                    p.hitlAt = datetime()
                """,
                pid=page_id,
            )
        return {"ok": True, "page_id": page_id, "status": "approved"}
    except Exception as exc:
        log.exception("Approve failed for %s: %s", page_id, exc)
        raise HTTPException(status_code=500, detail=str(exc))


@router.post("/flag")
async def flag_page(
    payload: dict = Body(...),
    driver: Driver = Depends(get_driver),
) -> dict[str, Any]:
    page_id = payload.get("page_id")
    note = payload.get("note", "")
    if not page_id:
        raise HTTPException(status_code=422, detail="page_id required")
    try:
        with driver.session() as s:
            s.run(
                """
                MATCH (p:PAGE {id: $pid})
                SET p.hitlStatus = 'flagged',
                    p.hitlNote = $note,
                    p.hitlAt = datetime()
                """,
                pid=page_id,
                note=note,
            )
        return {"ok": True, "page_id": page_id, "status": "flagged"}
    except Exception as exc:
        raise HTTPException(status_code=500, detail=str(exc))


@router.post("/correct")
async def submit_correction(
    payload: dict = Body(...),
    driver: Driver = Depends(get_driver),
) -> dict[str, Any]:
    page_id = payload.get("page_id")
    chunk_id = payload.get("chunk_id")
    corrected_text = payload.get("corrected_text", "")
    problem_class = payload.get("problem_class", "")
    notes = payload.get("notes", "")

    if not page_id:
        raise HTTPException(status_code=422, detail="page_id required")

    try:
        # Use existing correction_writer if chunk_id supplied
        if chunk_id and corrected_text:
            # Fetch the original text to populate CorrectionInput.before
            before = ""
            try:
                with driver.session() as s:
                    row = s.run("MATCH (c:CHUNK {id: $id}) RETURN c.text AS t", id=chunk_id).single()
                    before = (row["t"] or "") if row else ""
            except Exception:
                pass
            try:
                write_correction(
                    driver,
                    CorrectionInput(
                        page_id=page_id,
                        before=before,
                        after=corrected_text,
                        problem_class=problem_class or "OK",
                        editor_id="hitl-ui",
                    ),
                )
            except Exception as exc:
                log.warning("correction_writer failed, falling back to direct write: %s", exc)
                with driver.session() as s:
                    s.run(
                        """
                        MATCH (c:CHUNK {id: $cid})
                        SET c.hitlCorrectedText = $text,
                            c.hitlProblemClass = $pc,
                            c.hitlNotes = $notes,
                            c.hitlAt = datetime()
                        """,
                        cid=chunk_id,
                        text=corrected_text,
                        pc=problem_class,
                        notes=notes,
                    )
        # Mark page as corrected
        with driver.session() as s:
            s.run(
                """
                MATCH (p:PAGE {id: $pid})
                SET p.hitlStatus = 'corrected',
                    p.hitlAt = datetime()
                """,
                pid=page_id,
            )
        return {"ok": True, "page_id": page_id, "chunk_id": chunk_id}
    except Exception as exc:
        log.exception("Correction failed: %s", exc)
        raise HTTPException(status_code=500, detail=str(exc))


# ---------------------------------------------------------------------------
# Skip (move to end of queue)
# ---------------------------------------------------------------------------


@router.post("/skip")
async def skip_page(
    payload: dict = Body(...),
    driver: Driver = Depends(get_driver),
) -> dict[str, Any]:
    page_id = payload.get("page_id")
    if not page_id:
        raise HTTPException(status_code=422, detail="page_id required")
    with driver.session() as s:
        s.run(
            "MATCH (p:PAGE {id: $pid}) SET p.hitlSkipped = true, p.hitlAt = datetime()",
            pid=page_id,
        )
    return {"ok": True, "page_id": page_id}


# ---------------------------------------------------------------------------
# OCR HITL fusion review (Track E, Feature 4) — 4-column compare + fuse choice
# ---------------------------------------------------------------------------


@router.get("/ocr/{page_id}")
async def ocr_compare(page_id: str, driver: Driver = Depends(get_driver)) -> dict[str, Any]:
    """Return all engine texts + current fused text for the 4-column compare.

    Columns: original image | Paddle | LLM(s) (Qwen/DeepSeek + Custom when
    present) | preference. The frontend renders the original image via
    ``/api/image/{page_id}?variant=original``.
    """
    with driver.session() as s:
        row = s.run(
            """
            MATCH (p:PAGE {id: $pid})
            RETURN p.id AS page_id, p.documentId AS document_id,
                   p.docPageIndex AS page_index, p.language AS language, p.tier AS tier,
                   p.imageUri AS image_uri,
                   p.preprocessedImageUri AS preprocessed_uri,
                   p.paddleOcrText AS paddle, p.paddleOcrConfidence AS paddle_conf,
                   p.qwenVlOcrText AS qwen, p.qwenVlOcrConfidence AS qwen_conf,
                   p.deepseekOcrText AS deepseek, p.deepseekOcrConfidence AS deepseek_conf,
                   p.customOcrText AS custom, p.customOcrConfidence AS custom_conf,
                   p.customOcrModelVersion AS custom_model,
                   p.textFused AS fused, p.fusionSource AS fusion_source,
                   p.fusionWinnerLlm AS winner_llm, p.fusionAgreementRate AS agreement_rate
            """,
            pid=page_id,
        ).single()
    if row is None:
        raise HTTPException(status_code=404, detail=f"Page {page_id!r} not found")

    d = dict(row)
    engines = []
    for key, label in (("paddle", "PaddleOCR"), ("qwen", "Qwen-VL"),
                       ("deepseek", "DeepSeek"), ("custom", d.get("custom_model") or "Custom")):
        text = d.get(key)
        if text is not None:
            engines.append({
                "key": key,
                "label": label,
                "text": text,
                "confidence": d.get(f"{key}_conf"),
                "char_count": len(text or ""),
            })

    # Only advertise an image_url when the page actually has a stored image.
    # Pages that came from a non-rasterised native PDF or whose upload pipeline
    # never reached the rasterise stage have imageUri=NULL; emitting a URL
    # for those would just give the browser a 404 and a broken-image icon.
    has_original = bool(d.get("image_uri"))
    has_preprocessed = bool(d.get("preprocessed_uri"))
    image_url = (
        f"/api/image/{page_id}?variant=original" if has_original else None
    )
    preprocessed_url = (
        f"/api/image/{page_id}?variant=preprocessed" if has_preprocessed else None
    )

    return {
        "page_id": page_id,
        "document_id": d["document_id"],
        "page_index": d["page_index"],
        "language": d["language"],
        "tier": d["tier"],
        "image_url": image_url,
        "preprocessed_image_url": preprocessed_url,
        "image_available": has_original or has_preprocessed,
        "engines": engines,
        "fused": d["fused"],
        "fusion_source": d["fusion_source"],
        "winner_llm": d["winner_llm"],
        "agreement_rate": d["agreement_rate"],
    }


@router.post("/fuse")
async def hitl_fuse(payload: dict = Body(...), driver: Driver = Depends(get_driver)) -> dict[str, Any]:
    """Apply a human fusion decision.

    Body: ``{page_id, choice, manual_text?}`` where ``choice`` is one of
    ``'paddle' | 'qwen' | 'deepseek' | 'custom' | 'auto' | 'manual'``.

    - engine key -> copy that engine's text to ``textFused``;
    - ``'manual'`` -> store the user-edited ``manual_text``;
    - ``'auto'``  -> re-run the deterministic 3-engine fusion.
    """
    page_id = payload.get("page_id")
    choice = payload.get("choice", "manual")
    manual_text = payload.get("manual_text", "")
    if not page_id:
        raise HTTPException(status_code=422, detail="page_id required")

    with driver.session() as s:
        row = s.run(
            """
            MATCH (p:PAGE {id: $pid})
            RETURN p.paddleOcrText AS paddle, p.qwenVlOcrText AS qwen,
                   p.deepseekOcrText AS deepseek, p.customOcrText AS custom
            """,
            pid=page_id,
        ).single()
    if row is None:
        raise HTTPException(status_code=404, detail=f"Page {page_id!r} not found")

    engine_text = {
        "paddle": row["paddle"], "qwen": row["qwen"],
        "deepseek": row["deepseek"], "custom": row["custom"],
    }

    if choice == "manual":
        fused_text = manual_text
    elif choice == "auto":
        fused_text = _rerun_fusion(page_id, engine_text)
    elif choice in engine_text:
        fused_text = engine_text[choice] or ""
    else:
        raise HTTPException(status_code=400, detail=f"unknown choice {choice!r}")

    with driver.session() as s:
        s.run(
            """
            MATCH (p:PAGE {id: $pid})
            SET p.textFused = $text, p.fusionStatus = 'ok',
                p.fusionSource = 'hitl', p.fusionChoice = $choice,
                p.fusionCharCount = $n, p.hitlFusedAt = timestamp()
            """,
            pid=page_id, text=fused_text, choice=choice, n=len(fused_text or ""),
        )
    return {"ok": True, "page_id": page_id, "choice": choice, "char_count": len(fused_text or "")}


def _rerun_fusion(page_id: str, engine_text: dict[str, str | None]) -> str:
    """Reconstruct OCRPageResults from stored texts and re-run deterministic fusion."""
    from apps.backend.ocr.base import OCRPageResult
    from apps.backend.ocr.fusion import fuse_results

    def _mk(engine: str, text: str | None, conf: float) -> OCRPageResult | None:
        if not text:
            return None
        return OCRPageResult(
            engine=engine, model_version="hitl", page_id=page_id,
            text=text, confidence=conf, char_count=len(text),
        )

    paddle = _mk("paddleocr", engine_text.get("paddle"), 0.9) or OCRPageResult(
        engine="paddleocr", model_version="hitl", page_id=page_id, text="", confidence=0.0)
    qwen = _mk("qwen_vl_ocr", engine_text.get("qwen"), 0.82)
    deepseek = _mk("deepseek_ocr", engine_text.get("deepseek"), 0.85)
    result = fuse_results(paddle, deepseek, qwen)
    return result.text_fused


@router.post("/ocr-correct")
async def ocr_correct(payload: dict = Body(...), driver: Driver = Depends(get_driver)) -> dict[str, Any]:
    """Persist an OCR-level correction + a FEW_SHOT_EXAMPLE so the LLM 'learns'.

    Body: ``{page_id, before?, corrected_text, problem_class?, notes?}``.
    """
    page_id = payload.get("page_id")
    corrected = payload.get("corrected_text", "")
    before = payload.get("before", "")
    problem_class = payload.get("problem_class") or "OK"
    notes = payload.get("notes", "")
    if not page_id or not corrected:
        raise HTTPException(status_code=422, detail="page_id and corrected_text required")

    try:
        write_correction(
            driver,
            CorrectionInput(
                page_id=page_id, before=before, after=corrected,
                problem_class=problem_class, editor_id="hitl-ocr-ui",
            ),
        )
    except Exception as exc:  # noqa: BLE001
        log.warning("correction_writer failed: %s", exc)
    # Persist the note on the page for the per-document prompt memory.
    with driver.session() as s:
        s.run(
            """
            MATCH (p:PAGE {id: $pid})
            SET p.hitlOcrNote = $notes, p.hitlOcrCorrectedAt = timestamp()
            """,
            pid=page_id, notes=notes,
        )
    return {"ok": True, "page_id": page_id}


@router.get("/export/{page_id}")
async def export_page(
    page_id: str,
    format: str = Query("docx", description="docx | pdf"),
    driver: Driver = Depends(get_driver),
) -> Response:
    """Export a page's fused text as a Word (.docx) or PDF document."""
    with driver.session() as s:
        row = s.run(
            """
            MATCH (p:PAGE {id: $pid})
            OPTIONAL MATCH (d:DOCUMENT {id: p.documentId})
            RETURN p.textFused AS text, p.docPageIndex AS page_index,
                   coalesce(d.title, p.documentId) AS title
            """,
            pid=page_id,
        ).single()
    if row is None:
        raise HTTPException(status_code=404, detail=f"Page {page_id!r} not found")
    text = row["text"] or ""
    title = row["title"] or page_id
    page_index = row["page_index"]

    if format == "docx":
        data, media_type, ext = _export_docx(title, page_index, text), (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"), "docx"
    elif format == "pdf":
        data, media_type, ext = _export_pdf(title, page_index, text), "application/pdf", "pdf"
    else:
        raise HTTPException(status_code=400, detail="format must be docx or pdf")

    # page_id may contain CJK chars (cannot be latin-1 encoded in a raw header),
    # so emit an ASCII-safe fallback name plus an RFC 5987 UTF-8 filename*.
    ascii_name = page_id.encode("ascii", "ignore").decode("ascii").strip("._") or "page"
    quoted = quote(f"{page_id}.{ext}", safe="")
    disposition = (
        f'attachment; filename="{ascii_name}.{ext}"; '
        f"filename*=UTF-8''{quoted}"
    )
    return Response(
        content=data, media_type=media_type,
        headers={"Content-Disposition": disposition},
    )


def _export_docx(title: str, page_index: Any, text: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(str(title), level=1)
    if page_index is not None:
        doc.add_paragraph(f"Page {page_index}")
    for line in (text or "").split("\n"):
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _export_pdf(title: str, page_index: Any, text: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    # STSong-Light is a CID font bundled with reportlab — renders CJK without
    # shipping a TTF.
    font_name = "STSong-Light"
    try:
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    except Exception:  # noqa: BLE001 — fall back to Helvetica (latin-only)
        font_name = "Helvetica"

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=A4)
    styles = getSampleStyleSheet()
    body = ParagraphStyle("CJKBody", parent=styles["Normal"], fontName=font_name,
                          fontSize=12, leading=20)
    head = ParagraphStyle("CJKHead", parent=styles["Title"], fontName=font_name, fontSize=18)

    def _esc(s: str) -> str:
        return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story: list[Any] = [Paragraph(_esc(str(title)), head)]
    if page_index is not None:
        story.append(Paragraph(f"Page {page_index}", body))
    story.append(Spacer(1, 12))
    for line in (text or "").split("\n"):
        story.append(Paragraph(_esc(line) or "&nbsp;", body))
    pdf.build(story)
    return buf.getvalue()


@router.post("/ocr/approve")
async def ocr_approve(payload: dict = Body(...), driver: Driver = Depends(get_driver)) -> dict[str, Any]:
    """Approve a document's OCR/fusion and trigger the index pipeline.

    Body: ``{document_id}``. Sets ``DOCUMENT.status='approved'`` then spawns
    ``scripts/run_index_pipeline.py`` (chunk -> embed -> keywords) which sets
    ``status='indexed'`` when complete.
    """
    document_id = payload.get("document_id")
    if not document_id:
        raise HTTPException(status_code=422, detail="document_id required")

    with driver.session() as s:
        cur = s.run(
            "MATCH (d:DOCUMENT {id: $id}) RETURN d.status AS status",
            id=document_id,
        ).single()
        if cur is None:
            raise HTTPException(status_code=404, detail=f"Document {document_id!r} not found")
        status = cur["status"]
        # Only documents that have been through OCR review can be approved.
        # Block re-approval of an already-indexed doc (would re-spawn indexing).
        if status not in (None, "awaiting_review", "approved"):
            raise HTTPException(
                status_code=409,
                detail=f"document status is {status!r}; expected 'awaiting_review'",
            )
        s.run(
            "MATCH (d:DOCUMENT {id: $id}) SET d.status='approved', "
            "d.statusUpdatedAt=timestamp() RETURN d.id AS id",
            id=document_id,
        ).consume()

    try:
        cmd = [sys.executable, str(_INDEX_RUNNER), "--document-id", document_id]
        log_path = _REPO_ROOT / "staging" / f"index_{document_id[:16]}.log"
        log_path.parent.mkdir(parents=True, exist_ok=True)
        with open(log_path, "ab") as logf:
            subprocess.Popen(  # noqa: S603 — trusted args
                cmd, cwd=str(_REPO_ROOT), env=dict(os.environ),
                stdout=logf, stderr=subprocess.STDOUT, start_new_session=True,
            )
    except Exception as exc:  # noqa: BLE001
        log.exception("index pipeline spawn failed: %s", exc)
        raise HTTPException(status_code=500, detail=f"index spawn failed: {exc}")

    return {"ok": True, "document_id": document_id, "status": "approved", "indexing": "started"}
